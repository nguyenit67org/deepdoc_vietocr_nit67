"""Build the approved OCR demo as editable slides using the supplied VTS deck.

Run from the repository root with: .venv/bin/python data/reports/demo_slide_draft/build_vts_pptx.py
The linked public presentation provides its own master, decorative shapes and
Viettel Solutions logo; benchmark figures come from the full-inference bundle.
"""

from __future__ import annotations

import csv
from collections import Counter
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import requests
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
ASSET = ROOT / "assets"
OUTPUT = ROOT / "vts_ocr_demo_2026-09-29_report_v2.pptx"
METRICS_CSV = ROOT.parents[2] / "ocr-collab/artifacts/evidence_full_inference_all_v3/metrics.csv"
CAUSE_CSV = ROOT / "error_cause_full_inference_v3.csv"
TEMPLATE_URL = (
    "https://docs.google.com/presentation/d/"
    "1q_pOTwJijSkRhqOXbuSWdpNyFvmabVpO/export/pptx"
)

RED = RGBColor(0xEE, 0x00, 0x33)
INK = RGBColor(0x44, 0x49, 0x4D)
GRAY = RGBColor(0x73, 0x7A, 0x80)
LIGHT = RGBColor(0xF4, 0xF4, 0xF4)
BORDER = RGBColor(0xE3, 0xE3, 0xE3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x20, 0x70, 0x55)
TITLE_FONT = "Sarabun"
BODY_FONT = "Arial"


def clone_shape(source, target, shape):
    """Copy a real VTS decorative shape, including image-fill relationships."""
    element = deepcopy(shape.element)
    for blip in element.xpath(".//a:blip"):
        old_id = blip.get(qn("r:embed"))
        if old_id:
            image_part = source.part.related_part(old_id)
            blip.set(qn("r:embed"), target.part.relate_to(image_part, RT.IMAGE))
    target.shapes._spTree.insert_element_before(element, "p:extLst")


def base_slide(prs, source, indices):
    slide = prs.slides.add_slide(source.slide_layout)
    for idx in indices:
        clone_shape(source, slide, source.shapes[idx])
    return slide


def textbox(slide, x, y, w, h, text, *, size=20, color=INK, bold=False,
            font=BODY_FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
            margin=0.02, spacing=1.12):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    for i, line in enumerate(text.split("\n")):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.text = line
        para.alignment = align
        para.space_after = Pt(size * 0.14)
        para.line_spacing = spacing
        for run in para.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return shape


def panel(slide, x, y, w, h, *, fill=LIGHT, edge=BORDER, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = edge
    shape.line.width = Pt(0.6)
    return shape


def card(slide, x, y, w, h, title, body, *, tint=LIGHT, emphasis=None,
         title_size=20, body_size=18):
    panel(slide, x, y, w, h, fill=tint)
    if emphasis:
        textbox(slide, x + 0.25, y + 0.2, w - 0.5, 0.62, emphasis,
                size=32, color=RED, bold=True, font=TITLE_FONT)
        start = y + 0.94
    else:
        start = y + 0.28
    textbox(slide, x + 0.25, start, w - 0.5, 0.72, title,
            size=title_size, bold=True, font=TITLE_FONT)
    textbox(slide, x + 0.25, start + 0.75, w - 0.5,
            max(0.3, h - (start - y) - 1.0), body, size=body_size)


def badge(slide, x, y, w, text, *, fill=LIGHT, color=INK):
    panel(slide, x, y, w, 0.38, fill=fill, edge=fill)
    textbox(slide, x + 0.07, y + 0.025, w - 0.14, 0.30, text,
            size=10, color=color, bold=True, align=PP_ALIGN.CENTER)


def photo(slide, filename, x, y, w, h, *, caption="", bg=WHITE):
    panel(slide, x, y, w, h, fill=bg, radius=False)
    path = ASSET / filename
    with Image.open(path) as im:
        iw, ih = im.size
        data = BytesIO()
        # PowerPoint/Google Slides do not support WebP images inside PPTX.
        # Keep OCR crops lossless, and encode page thumbnails as quality JPEG.
        is_page = filename.startswith(("demo-", "stamp-", "sign-", "missing-"))
        im.save(data, "JPEG" if is_page else "PNG", quality=90)
        data.seek(0)
    image_h = h - (0.48 if caption else 0.12)
    scale = min((w - 0.15) / iw, (image_h - 0.13) / ih)
    pw, ph = iw * scale, ih * scale
    slide.shapes.add_picture(data, Inches(x + (w - pw) / 2),
                             Inches(y + (image_h - ph) / 2),
                             width=Inches(pw), height=Inches(ph))
    if caption:
        textbox(slide, x + 0.15, y + h - 0.42, w - 0.3, 0.34,
                caption, size=10, color=GRAY)


def quote(slide, x, y, w, label, value, *, good=False, h=0.95, size=17):
    fill = RGBColor(0xF2, 0xF7, 0xF4) if good else RGBColor(0xFA, 0xF0, 0xF0)
    panel(slide, x, y, w, h, fill=fill, edge=fill)
    textbox(slide, x + 0.17, y + 0.1, w - 0.34, 0.23, label.upper(),
            size=10, color=GREEN if good else RED, bold=True)
    textbox(slide, x + 0.17, y + 0.37, w - 0.34, h - 0.44, value,
            size=size, bold=True)


def banner(slide, y, text, *, size=21):
    panel(slide, 1.36, y, 17.4, 0.9,
          fill=RGBColor(0xFB, 0xF1, 0xF2), edge=RGBColor(0xF4, 0xD3, 0xD8))
    textbox(slide, 1.6, y + 0.15, 16.9, 0.6, text,
            size=size, color=INK, bold=True)


def metric_row(slide, x, y, w, name, key, metric):
    """Show measured F1(P,R)/NEM from the historical run, not its old status."""
    f1, precision, recall, nem = (float(metric[k]) for k in ("f1", "precision", "recall", "nem"))
    meets = f1 >= .85 and precision >= .85 and recall >= .85 and nem >= .75
    panel(slide, x, y, w, 0.89, fill=WHITE, edge=BORDER)
    textbox(slide, x + .15, y + .11, w * .48, .65, f"{name} ({key})",
            size=15, bold=True)
    value = f"{f1:.3f} ({precision:.3f}, {recall:.3f}) / {nem:.3f}".replace(".", ",")
    textbox(slide, x + w * .51, y + .17, w * .46, .52, value,
            size=15, color=GREEN if meets else RED, bold=True,
            align=PP_ALIGN.RIGHT)


def white_logo_png(template):
    # Recolor the template's own transparent Viettel Solutions logo for red covers.
    from zipfile import ZipFile

    with ZipFile(BytesIO(template)) as archive:
        img = Image.open(BytesIO(archive.read("ppt/media/image3.png"))).convert("RGBA")
    pixels = img.load()
    for yy in range(img.height):
        for xx in range(img.width):
            r, g, b, a = pixels[xx, yy]
            if a:
                pixels[xx, yy] = (255, 255, 255, a)
    stream = BytesIO()
    img.save(stream, "PNG")
    stream.seek(0)
    return stream


def make_deck():
    cache = Path("/tmp/opencode/vts_template_source.pptx")
    if cache.exists():
        template = cache.read_bytes()
    else:
        response = requests.get(TEMPLATE_URL, timeout=90)
        response.raise_for_status()
        template = response.content
    if not template.startswith(b"PK"):
        raise ValueError("VTS template download is not a PowerPoint file")
    prs = Presentation(BytesIO(template))
    cover_source, content_source, closing_source = prs.slides[0], prs.slides[3], prs.slides[23]
    original_count = len(prs.slides)
    white_logo = white_logo_png(template)
    with METRICS_CSV.open(newline="", encoding="utf-8") as source:
        metrics_by_key = {row["field_name"]: row for row in csv.DictReader(source)}
    with CAUSE_CSV.open(newline="", encoding="utf-8") as source:
        cause_rows = list(csv.DictReader(source))
    assert len(cause_rows) == 229
    causes = Counter(row["report_group"] for row in cause_rows)
    pages = []

    def content(title, source, section, *, number=None):
        print(f"Composing {len(pages) + 1:02d}: {title}", flush=True)
        slide = base_slide(prs, content_source, range(8))
        pages.append(slide)
        n = number or len(pages)
        textbox(slide, 1.33, 0.59, 15.5, 0.68, title,
                size=28, color=INK, bold=True, font=TITLE_FONT)
        textbox(slide, 1.35, 10.27, 14.3, 0.35, source,
                size=9, color=GRAY)
        textbox(slide, 18.78, 10.27, 0.65, 0.35, str(n).zfill(2),
                size=10, color=RED, align=PP_ALIGN.RIGHT)
        badge(slide, 1.36, 1.43, min(5, 1.2 + len(section) * .11), section.upper(),
              fill=RGBColor(0xF8, 0xE7, 0xEA), color=RED)
        return slide

    # 01 — cover, using actual VTS cover artwork rather than a lookalike.
    print("Composing 01: Cover", flush=True)
    s = base_slide(prs, cover_source, range(6)); pages.append(s)
    s.shapes.add_picture(white_logo, Inches(16.25), Inches(0.75), width=Inches(3.0))
    textbox(s, 8.05, 3.9, 10.8, 1.7, "AI OCR cho văn bản\nhành chính",
            size=44, color=WHITE, bold=True, font=TITLE_FONT)
    textbox(s, 8.08, 6.15, 10.6, 1.1,
            "Demo pipeline 01 · Văn phòng số / VPS", size=23, color=WHITE)
    textbox(s, 8.08, 7.43, 10.6, 0.65,
            "Báo cáo 29/09/2026  |  Full inference: 35/100 tài liệu đạt",
            size=16, color=WHITE)

    # 02 — context.
    s = content("Từ nhập liệu thủ công đến dữ liệu có cấu trúc",
                "Outline mục 1 · bài toán Văn phòng số", "01 · Bối cảnh")
    card(s, 1.35, 2.06, 5.5, 5.7, "Hiện trạng",
         "Nhân viên đọc PDF rồi nhập lại số ký hiệu, ngày, cơ quan, trích yếu… vào hệ thống.", emphasis="01")
    card(s, 7.07, 2.06, 5.5, 5.7, "Pain point",
         "Khối lượng xử lý lớn kéo theo thời gian, nhân lực kiểm tra và nguy cơ sai sót.", emphasis="02")
    card(s, 12.78, 2.06, 5.9, 5.7, "Mục tiêu PoC",
         "Đánh giá OCR/KIE in-house có thể tùy chỉnh theo nghiệp vụ và thử nghiệm tại UAT VPS.", emphasis="03")
    banner(s, 8.15, "Đây là bước đánh giá khả thi, chưa phải cam kết thay thế API hiện tại.")

    # 03 — Agile approach.
    s = content("Viettel Agile: học từ dữ liệu thật", "Outline mục 1; benchmark policy §6, §9", "01 · Bối cảnh")
    steps = [("Aware", "Nhận diện pain point"), ("Go Through", "Khảo sát nghiệp vụ"),
             ("Implement", "Xây baseline"), ("Learn", "Đo và phân tích lỗi"),
             ("Enable", "Hướng tới áp dụng")]
    for i, (step, detail) in enumerate(steps):
        x = 1.36 + i * 3.53
        card(s, x, 2.32, 3.3, 2.95, step, detail,
             tint=RGBColor(0xFA, 0xF1, 0xF3) if i == 3 else LIGHT,
             title_size=20, body_size=16)
    card(s, 1.36, 5.72, 8.45, 3.38, "Xuất phát từ nghiệp vụ",
         "Đối chiếu với biểu mẫu/BA: trường nào OCR đọc, trường nào người dùng nhập, trường nào hệ thống tự điền?",
         body_size=19)
    card(s, 10.1, 5.72, 8.67, 3.38, "Cải tiến có đo lường",
         "Mọi thay đổi cần GT, benchmark, trace và regression test; không đặc biệt hóa theo tên file hoặc đáp án.",
         body_size=19)

    # 04 — dataset.
    s = content("Benchmark 100 PDF: đủ đa dạng để thấy điểm yếu",
                "Full inference manifest.json + predictions.jsonl", "02 · Phạm vi")
    for i, (value, title, detail) in enumerate([
        ("100", "văn bản", "Có GT theo tài liệu"),
        ("2.249", "trang", "1–353 trang/file"),
        ("69", "bản scan", "Chất lượng OCR biến thiên"),
        ("31", "bản điện tử", "Print/digital PDF"),
    ]):
        card(s, 1.36 + 4.42 * i, 2.02, 4.13, 3.35, title, detail,
             emphasis=value, title_size=20, body_size=16)
    card(s, 1.36, 5.85, 8.45, 3.38, "Loại văn bản",
         "Quyết định (30), Công văn (17), Nghị quyết, Thông báo, Thông tư, Kế hoạch…", body_size=19)
    card(s, 10.1, 5.85, 8.67, 3.38, "Nguồn / chủ thể",
         "Chính phủ, bộ/ngành, địa phương, sở giao dịch chứng khoán và doanh nghiệp; chưa phủ đủ lớp hiếm.", body_size=19)

    # 05 — contract.
    s = content("Định nghĩa đầu ra trước khi đo chất lượng",
                "docs/ocr_pipeline_benchmark_metrics.md · §1, §5", "02 · Phạm vi")
    card(s, 1.36, 2.03, 8.45, 5.66, "7 trường required",
         "type · title · code · documentDate\nofficeSender · priority_level · security_level\n\n5 trường khác được chấm:\nrecipients · signer · first_recipients\nsigner_title · province",
         body_size=19)
    card(s, 10.1, 2.03, 8.67, 5.66, "API 13 key; benchmark 12 field",
         "Theo chuẩn hiện hành, receiverDate phải trả chuỗi rỗng và không chấm. Bundle lịch sử còn lưu giá trị legacy; bản báo cáo loại field này.\n\nDocument Success@5%: tất cả required field có GT trong một tài liệu đều đạt bucket hợp lệ.",
         tint=RGBColor(0xFA, 0xF1, 0xF3), body_size=18)
    banner(s, 8.02, "Exact / Acceptable / Wrong / Missing / Hallucination / Correct Absent.", size=20)

    # Field names in Vietnamese first, API keys in parentheses, grouped by policy.
    s = content("Danh sách field theo nhóm nghiệp vụ",
                "docs/ocr_pipeline_benchmark_metrics.md · §1 · API 13 key / chấm 12 field",
                "02 · Required / critical / optional")
    card(s, 1.36, 2.04, 8.18, 7.3, "Critical + required · 7",
         "Loại văn bản (type)\nTrích yếu/tiêu đề (title)\nSố, ký hiệu văn bản (code)\nNgày ban hành (documentDate)\nCơ quan ban hành (officeSender)\nĐộ khẩn (priority_level)\nĐộ mật (security_level)",
         body_size=18)
    card(s, 9.81, 2.04, 5.17, 7.3, "Critical + optional · 4",
         "Nơi nhận (recipients)\nNgười ký (signer)\nKính gửi (first_recipients)\nChức danh người ký (signer_title)",
         body_size=16)
    card(s, 15.23, 2.04, 3.55, 3.42, "Non-critical + optional",
         "Địa danh (province)", title_size=17, body_size=17)
    card(s, 15.23, 5.71, 3.55, 3.63, "Out of scope",
         "Ngày tiếp nhận (receiverDate)\nGiữ key, không chấm.",
         title_size=17, body_size=16)

    # 06 — process.
    s = content("Team đã xây gì, và pipeline chạy ra sao?",
                "Outline mục 2; benchmark policy §6–§8", "02 · Phạm vi")
    for i, text in enumerate(["PDF", "Layout + OCR", "Extraction", "API chuẩn hóa", "GT + benchmark"]):
        panel(s, 1.36 + i * 3.52, 2.26, 3.27, 1.45, fill=LIGHT)
        textbox(s, 1.54 + i * 3.52, 2.65, 2.93, 0.65, text,
                size=21, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
    for i, (title, body) in enumerate([
        ("Chuẩn bị", "Thu thập dữ liệu; làm GT; thống nhất quy tắc với BA."),
        ("Thử nghiệm", "Khảo sát pipeline riêng; chọn pipeline 01 để demo."),
        ("Truy vết", "Đo accuracy, thời gian và RAM/GPU; giữ OCR, bbox, ảnh trang."),
    ]):
        card(s, 1.36 + i * 5.83, 4.38, 5.55, 3.5, title, body, body_size=18)
    banner(s, 8.3, "Số liệu trình bày lấy từ full inference, không trộn tài nguyên của replay.")

    # 07 — demo guide.
    s = content("Demo theo 3 mức độ: dễ → phổ biến → khó",
                "Full inference predictions.jsonl + document_success.csv", "03 · Demo")
    card(s, 1.36, 2.08, 5.5, 5.58, "Case dễ",
         "3116+QĐ-BKHCN; 26+2026+NQ-HĐND.\nKhông có Wrong/Missing/Hallucination.",
         emphasis="01", body_size=19)
    card(s, 7.08, 2.08, 5.5, 5.58, "Case phổ biến",
         "Quyết định 1666 sạch; Công văn 8906 đạt 7/7 required nhưng còn 2 Missing optional.",
         emphasis="02", body_size=19)
    card(s, 12.8, 2.08, 5.98, 5.58, "Case khó",
         "8931, 8948, 29+CT: scan; lỗi ngày, mức khẩn, chữ ký và chức danh.",
         emphasis="03", body_size=19)
    banner(s, 8.12, "PDF → output → GT/prediction → ảnh và trace → thời gian xử lý.")

    # 08 — easy cases.
    s = content("Văn bản rõ: trích xuất đúng và truy vết được",
                "Full inference predictions.jsonl + field_errors.csv", "03 · Demo / Case dễ")
    photo(s, "demo-3116.webp", 1.37, 2.02, 7.15, 7.7, caption="3116+QĐ-BKHCN · trace trang 1")
    card(s, 8.86, 2.02, 9.9, 3.54, "3116+QĐ-BKHCN · 2 trang",
         "4,37 s/file · PASS. Theo 12 field hiện hành: 10 Exact, title Acceptable, first_recipients Correct Absent.",
         body_size=19)
    card(s, 8.86, 5.82, 9.9, 3.42, "26+2026+NQ-HĐND · 5 trang",
         "7,93 s/file · PASS; không có bucket Acceptable, Wrong, Missing hoặc Hallucination.", body_size=19)

    # 09 — common cases.
    s = content("Quyết định và Công văn chiếm 47% tập mẫu",
                "Full inference predictions.jsonl + field_errors.csv", "03 · Demo / Phổ biến")
    photo(s, "demo-1666.webp", 1.36, 2.07, 5.6, 6.75, caption="1666+QĐ-TTg · scan 5 trang")
    photo(s, "demo-8906.webp", 7.16, 2.07, 5.6, 6.75, caption="8906+VPCP-CN · print 2 trang")
    card(s, 13.02, 2.07, 5.75, 3.16, "Quyết định 1666",
         "12,52 s/file; không field nào Wrong/Missing/Hallucination.", body_size=17)
    card(s, 13.02, 5.49, 5.75, 3.33, "Công văn 8906",
         "4,89 s/file; 7/7 required đạt. recipients và first_recipients Missing.", body_size=17)

    # 10 — difficult cases.
    s = content("Header đúng chưa đồng nghĩa đúng cả văn bản",
                "Full inference field_errors.csv + document_success.csv", "03 · Demo / Khó")
    card(s, 1.36, 2.16, 5.52, 5.72, "8931+VPCP-PL",
         "4,42 s · 1 trang scan. Ngày 29→19; cơ quan dính chữ; Hỏa tốc thành Bình thường. FAIL 3 required.",
         body_size=19)
    card(s, 7.1, 2.16, 5.52, 5.72, "8948+VPCP-CN",
         "6,50 s · 2 trang scan. Mất ngày, mất Kính gửi; vùng chức danh lẫn OCR rác. FAIL vì ngày.",
         body_size=19)
    card(s, 12.84, 2.16, 5.93, 5.72, "29+CT-TTg",
         "7,76 s · 4 trang scan. Header chính Exact; rớt Hỏa tốc, tên ký sai, chức danh Missing.",
         body_size=19)
    banner(s, 8.32, "Các slide lỗi sau đây gắn GT/prediction với đúng ảnh và OCR trong trace.")

    # 11 — results and provenance.
    s = content("35/100 văn bản đạt; PoC còn cách 45 văn bản",
                "summary.json · full inference v3 · metric v2.1", "04 · Kết quả")
    card(s, 1.36, 2.03, 5.55, 3.22, "Document Success@5%",
         "35 PASS / 100. PoC ≥80%; stretch ≥90%.", emphasis="35%", body_size=18)
    card(s, 7.13, 2.03, 5.55, 3.22, "Thời gian trung bình",
         "44,1 s/file (wall-clock, 100 PDF).", emphasis="44,1 s", body_size=18)
    card(s, 12.9, 2.03, 5.88, 3.22, "Giây/trang*",
         "1,14–5,85 s; *TB của giây/file ÷ trang/file.", emphasis="2,48 s", body_size=18)
    card(s, 1.36, 5.62, 8.45, 3.53, "RAM tiến trình",
         "Trung bình 2.745 MB · đỉnh 6.012 MB.\nBenchmark chạy trên Quadro RTX 5000.", body_size=19)
    card(s, 10.1, 5.62, 8.68, 3.53, "VRAM: phân biệt phạm vi đo",
         "Device VRAM toàn GPU: 3.077 / 6.793 MB (TB/đỉnh). Process VRAM: chưa đo được; không gán 0 hay dùng allocator để thay.", body_size=18)

    # 12 — field results, not an error chart.
    s = content("Điểm nghẽn: chức danh ký, trích yếu, ngày",
                "Full inference metrics.csv · receiverDate ngoài phạm vi hiện hành", "04 · Kết quả theo field")
    card(s, 1.36, 2.04, 8.46, 6.92, "Tín hiệu tích cực",
         "province: F1 = 0,984 — field duy nhất đạt đủ target.\n\ntype, security_level: F1 = 1,00 nhưng chưa đánh giá đủ lớp hiếm/mẫu GT trống.\n\nKhông có Hallucination trên những field có GT trống trong bundle này.",
         body_size=20)
    card(s, 10.1, 2.04, 8.68, 6.92, "Field cần ưu tiên",
         "signer_title  F1 0,380  ·  51 Wrong + 12 Missing\n\ntitle  F1 0,670  ·  31 Wrong + 3 Missing\n\nofficeSender  F1 0,735  ·  24 Wrong + 4 Missing\n\ndocumentDate  F1 0,785  ·  10 Wrong + 19 Missing",
         tint=RGBColor(0xFA, 0xF1, 0xF3), body_size=18)

    # Two legible tables rather than a tiny all-field chart. Raw metrics are
    # authoritative; stored status/target columns in this older bundle are not.
    s = content("Target và kết quả field: Critical + required",
                "Full inference v3 · metrics.csv; target báo cáo hiện hành, không dùng status cũ",
                "04 · F1 (Precision, Recall) / NEM")
    textbox(s, 1.43, 2.04, 17.1, .56,
            "Target mỗi field: F1 (P, R) / NEM ≥ 0,85 (0,85, 0,85) / 0,75",
            size=21, color=RED, bold=True)
    required = [
        ("Loại văn bản", "type"), ("Trích yếu", "title"),
        ("Số, ký hiệu", "code"), ("Ngày ban hành", "documentDate"),
        ("Cơ quan ban hành", "officeSender"), ("Độ khẩn", "priority_level"),
        ("Độ mật", "security_level"),
    ]
    for i, (name, key) in enumerate(required):
        metric_row(s, 1.36, 2.85 + .96 * i, 17.42, name, key, metrics_by_key[key])
    textbox(s, 1.42, 9.6, 17.05, .44,
            "F1 = 1,000 ở type/security_level chưa đủ phủ lớp hiếm hoặc mẫu GT trống; màu chỉ đối chiếu số học.",
            size=11, color=GRAY)

    s = content("Target và kết quả field: Critical / Non-critical optional",
                "Full inference v3 · metrics.csv; 12 field được chấm, receiverDate ngoài phạm vi",
                "04 · F1 (Precision, Recall) / NEM")
    textbox(s, 1.43, 2.04, 17.1, .56,
            "Target mỗi field: F1 (P, R) / NEM ≥ 0,85 (0,85, 0,85) / 0,75",
            size=21, color=RED, bold=True)
    textbox(s, 1.42, 2.76, 17.0, .4, "CRITICAL + OPTIONAL · 4 FIELD",
            size=16, color=RED, bold=True)
    optional = [
        ("Nơi nhận", "recipients"), ("Người ký", "signer"),
        ("Kính gửi", "first_recipients"), ("Chức danh người ký", "signer_title"),
    ]
    for i, (name, key) in enumerate(optional):
        metric_row(s, 1.36, 3.22 + .99 * i, 17.42, name, key, metrics_by_key[key])
    textbox(s, 1.42, 7.31, 17, .42, "NON-CRITICAL + OPTIONAL · 1 FIELD",
            size=16, color=RED, bold=True)
    metric_row(s, 1.36, 7.81, 17.42, "Địa danh", "province", metrics_by_key["province"])
    textbox(s, 1.42, 9.1, 17.12, .72,
            "Hallucination rate giữ ngưỡng cũ: ≤ 0,01; riêng Địa danh (province) ≤ 0,02. Mẫu số thiếu → chưa đánh giá đầy đủ.",
            size=15, color=GRAY)

    # 13 — OCR evidence.
    s = content("OCR: một ký tự sai làm hỏng field required",
                "8931+VPCP-PL · trace v2 trang 1 · issuance-date / upper-left-organization",
                "04 · Evidence / OCR")
    photo(s, "ocr-date-8931.webp", 1.36, 2.14, 7.65, 2.2,
          caption="Ngày · bbox [847, 274, 1337, 308]")
    photo(s, "ocr-sender-8931.webp", 1.36, 4.7, 7.65, 2.2,
          caption="Cơ quan · bbox [236, 154, 663, 195]")
    quote(s, 9.35, 2.14, 9.43, "GT · documentDate", "29/08/2026", good=True)
    quote(s, 9.35, 3.31, 9.43, "Pred / OCR", "19/08/2026 · ngày 19 tháng 8 năm 2026")
    quote(s, 9.35, 4.67, 9.43, "GT · officeSender", "VĂN PHÒNG CHÍNH PHỦ", good=True)
    quote(s, 9.35, 5.84, 9.43, "Pred / OCR", "VĂN PHÒNG CHÍNH PHỦNH PHỦ")
    textbox(s, 9.37, 7.13, 9.32, 2.2,
            "Extraction bám OCR nguồn. Metadata có chữ viết tay nhưng chưa đủ evidence để quy hai lỗi này cụ thể cho chữ viết tay.",
            size=18)

    # 14 — stamp/signature.
    s = content("Dấu mộc / vùng ký: header đúng nhưng mức khẩn rớt",
                "29+CT-TTg · trace v2 trang 1 & 4 · signature-name",
                "04 · Evidence / Dấu & ký")
    photo(s, "stamp-29.webp", 1.36, 2.07, 4.98, 7.52,
          caption="Trang 1 · kiểm tra dấu Hỏa tốc")
    photo(s, "sign-29.webp", 6.54, 2.07, 4.98, 7.52,
          caption="Trang 4 · OCR 'Huốc Dũng'")
    quote(s, 11.79, 2.07, 6.98, "GT / Pred · priority_level",
          "HỎA TỐC → BÌNH THƯỜNG", h=1.35, size=16)
    quote(s, 11.79, 3.62, 6.98, "GT / Pred · signer",
          "Hồ Quốc Dũng → Huốc Dũng", h=1.35, size=16)
    quote(s, 11.79, 5.17, 6.98, "GT / Pred · signer_title",
          "KT. THỦ TƯỚNG PHÓ THỦ TƯỚNG → rỗng", h=1.38, size=14)
    textbox(s, 11.85, 6.88, 6.9, 2.4,
            "priority_level không có evidence. Cần review ảnh màu vùng dấu; chưa kết luận dấu gây ra tất cả các lỗi vùng ký.", size=17)

    # 15 — layout / crop.
    s = content("Layout / crop: vùng text lẫn nội dung",
                "12+NQ-PVIRe và 240+TB-PVIRe · trace v2 trang 1",
                "04 · Evidence / Layout")
    photo(s, "layout-title-12.webp", 1.36, 2.13, 8.42, 4.05,
          caption="12+NQ · khung đỏ = bbox prediction evidence")
    photo(s, "layout-sender-240.webp", 10.1, 2.13, 8.68, 4.05,
          caption="240+TB · bbox của 'Số:' nhận nhầm")
    card(s, 1.36, 6.45, 8.42, 2.73, "title · Wrong",
         "GT: 'Về việc chốt danh sách cổ đông…'\nPred: thêm tổ chức và HỘI ĐỒNG QUẢN TRỊ.",
         body_size=16)
    card(s, 10.1, 6.45, 8.68, 2.73, "officeSender · Wrong",
         "GT: TỔNG CÔNG TY CỔ PHẦN TÁI BẢO HIỂM PVI\nPred: 'Số:' — bắt nhầm anchor.", body_size=16)

    # 16 — extraction.
    s = content("Extraction: OCR có text, hậu xử lý giữ nhầm stub",
                "2125+TB-ĐLTKV · trace v2 trang 1 · kinh-gui-anchor",
                "04 · Evidence / Extraction")
    photo(s, "extract-2125.webp", 1.36, 2.14, 17.4, 2.83,
          caption="Kính gửi · bbox [473, 616, 1450, 767]")
    quote(s, 1.36, 5.25, 17.4, "GT · first_recipients",
          "Ủy ban Chứng khoán Nhà nước; Sở Giao dịch Chứng khoán Hà Nội; Trung tâm Lưu ký…",
          good=True, h=1.1, size=17)
    quote(s, 1.36, 6.55, 17.4, "Prediction",
          "Uỷ ban Chứng khoán Nhà nước; số 1.772; Sở Giao dịch Chứng khoán Hà Nội; Trung tâm…",
          h=1.15, size=17)
    textbox(s, 1.45, 8.12, 17.05, 1.06,
            "Thử lọc stub khi tách danh sách; regression test để không xóa số thật trong tên người nhận.",
            size=19)

    # 17 — missing.
    s = content("Missing: không thấy evidence ≠ biết vị trí GT",
                "8948+VPCP-CN · trace v2 trang 1 · Missing không có evidence",
                "04 · Evidence / Missing")
    photo(s, "missing-8948.webp", 1.36, 2.06, 7.12, 7.53,
          caption="Trang 1 do người review chọn; chưa là bbox GT")
    quote(s, 8.78, 2.06, 10, "GT / Pred · documentDate",
          "01/09/2026 → rỗng", h=1.1, size=18)
    quote(s, 8.78, 3.42, 10, "GT / Pred · first_recipients",
          "Bộ trưởng các Bộ…; Chủ tịch UBND… → rỗng", h=1.35, size=16)
    card(s, 8.78, 5.06, 10, 4.02, "Nhận định có giới hạn",
         "Cần soi layout/OCR và ranh giới văn bản trước khi thử fallback. Không suy vị trí GT từ trang tự chọn. Bundle không có Hallucination ở field có mẫu GT trống; không suy rộng sang lớp hiếm.",
         body_size=18)

    # 18 — GT review.
    s = content("Nghi vấn GT: bucket Wrong vẫn cần người review",
                "03+BC.HĐQT; 27+2026+VBHN · full inference trace v2",
                "04 · Evidence / GT")
    photo(s, "gt-03.webp", 1.36, 2.1, 8.43, 3.35,
          caption="OCR nguồn: Số 03/BC HĐQT")
    photo(s, "gt-27.webp", 10.1, 2.1, 8.68, 3.35,
          caption="OCR nguồn: Số 27/2026/VBHN-NĐ-BTC")
    card(s, 1.36, 5.7, 8.43, 3.58, "03+BC.HĐQT",
         "GT 03/BC.HĐQT · Pred 03/BC-HĐQT. Quy tắc API đổi dấu cách thành '-'; cần rà lại policy gán dấu.",
         body_size=18)
    card(s, 10.1, 5.7, 8.68, 3.58, "27+2026+VBHN",
         "GT …-ND-BTC · Pred …-NĐ-BTC. Ảnh/trace ủng hộ Đ; cần reviewer xác nhận, không tự sửa GT.",
         body_size=18)

    # Auditable aggregate distribution from the same 35/100 full-inference
    # bundle. This deliberately does not claim a fully reviewed root cause.
    s = content("Phân bố nhóm nguyên nhân lỗi · sơ bộ",
                "error_cause_full_inference_v3.csv · 229 ô lỗi / 12 field · full inference v3",
                "04 · Thống kê nguyên nhân")
    textbox(s, 1.43, 2.05, 17.1, .67,
            "n = 229 ô lỗi: 156 Wrong + 73 Missing; không tính receiverDate.",
            size=21, bold=True, color=INK)
    for i, (name, color) in enumerate([
        ("OCR/chính tả", RED), ("Trích xuất", INK),
        ("Dấu mộc", RGBColor(0xBE, 0x50, 0x63)), ("Khác", RGBColor(0xA6, 0xAA, 0xAE)),
    ]):
        y = 3.02 + 1.39 * i
        value = causes[name]
        textbox(s, 1.44, y - .06, 4.0, .55, name, size=21, bold=True)
        panel(s, 5.44, y + .03, 10.44, .59, fill=LIGHT, edge=LIGHT, radius=False)
        panel(s, 5.44, y + .03, 10.44 * value / 109, .59,
              fill=color, edge=color, radius=False)
        textbox(s, 16.16, y - .01, 2.6, .63,
                f"{value}  ({100 * value / 229:.1f}%)".replace(".", ","),
                size=20, bold=True, color=color, align=PP_ALIGN.RIGHT)
    panel(s, 1.36, 8.84, 17.42, 1.03, fill=RGBColor(0xFA, 0xF1, 0xF3))
    textbox(s, 1.63, 8.96, 16.87, .78,
            "216 ô gán nhãn heuristic; 8 case evidence override + 5 dấu HỎA TỐC đối chiếu ảnh. "
            "Chưa review toàn bộ; Khác gồm thiếu evidence/nghi vấn GT.",
            size=14, color=INK)

    # 19 — feasibility.
    s = content("Khả thi để tiếp tục R&D; chưa đủ điều kiện thay thế",
                "Outline mục 5; full inference summary.json", "05 · Đánh giá")
    card(s, 1.36, 2.12, 8.44, 5.75, "Đã chứng minh",
         "Pipeline end-to-end trên 100 PDF; API dữ liệu cấu trúc.\n\nCó GT, benchmark, ảnh trang/OCR/bbox/trace và case print/scan chạy tốt.",
         body_size=20)
    card(s, 10.1, 2.12, 8.68, 5.75, "Chưa chứng minh",
         "35% Document Success còn cách PoC 80%.\n\nMẫu lớp hiếm/chụp nghiêng-ngang chưa đủ; chưa có so sánh chi phí/chất lượng cùng giao thức để kết luận thay thế.",
         tint=RGBColor(0xFA, 0xF1, 0xF3), body_size=19)
    banner(s, 8.17, "Cân bằng tổng quát hóa, chất lượng–GPU và hiểu ngữ cảnh–thời gian.", size=20)

    # 20 — next steps.
    s = content("Ưu tiên accuracy; sau đó tối ưu vận hành",
                "Outline mục 6; benchmark policy §6, §9", "06 · Tiếp theo")
    card(s, 1.36, 2.05, 5.5, 5.82, "01 · Dữ liệu",
         "Bổ sung hard cases dấu đỏ, chữ viết tay, trang nghiêng/ngang, lớp hiếm và văn bản dài; rà GT; chia train/dev/test nếu fine-tune.",
         body_size=18)
    card(s, 7.08, 2.05, 5.5, 5.82, "02 · Kỹ thuật",
         "Benchmark OCR riêng; thử model chữ viết tay, xử lý dấu/watermark, xác định ranh giới/trang cuối và chuẩn hóa giá trị hữu hạn.",
         body_size=18)
    card(s, 12.8, 2.05, 5.98, 5.82, "03 · Kiểm chứng",
         "Regression test mỗi vòng. Đạt accuracy target rồi tối ưu tốc độ → tài nguyên → chi phí → UAT có kiểm soát.",
         body_size=18)
    banner(s, 8.2, "Cần thống nhất mốc review, hard cases ưu tiên và nguồn lực GPU/chuyên môn.", size=20)

    # 21 — exact VTS closing artwork with our message and white VTS logo.
    s = base_slide(prs, closing_source, range(6)); pages.append(s)
    white_logo.seek(0)
    s.shapes.add_picture(white_logo, Inches(16.2), Inches(0.7), width=Inches(3.0))
    textbox(s, 7.77, 4.23, 11.1, 1.25, "Xin cảm ơn!",
            size=50, color=WHITE, bold=True, font=TITLE_FONT)
    textbox(s, 7.8, 5.68, 11.05, 1.5,
            "35/100 hôm nay · 80/100 là mốc PoC\nFeedback về field ưu tiên, UAT và nguồn lực?",
            size=23, color=WHITE)

    # Delete original lecture slides. The new slides retain the source deck's
    # slide master/theme and only the media relationships they actually use.
    for _ in range(original_count):
        slide_id = prs.slides._sldIdLst[0]
        prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)
    assert len(prs.slides) == len(pages) == 25
    prs.core_properties.title = "Demo AI OCR văn bản hành chính · Viettel Solutions"
    prs.core_properties.subject = "Pipeline 01 · full inference 35/100 · 29/09/2026"
    prs.save(OUTPUT)
    print(f"Saved {OUTPUT} ({OUTPUT.stat().st_size:,} bytes; {len(prs.slides)} slides)")


if __name__ == "__main__":
    make_deck()
