"""Generate vts_ocr_demo_antigravity.pptx using the official Viettel Solutions template.

This script creates the standalone presentation for Google Slides upload
without modifying any existing files.
"""

from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import os

import requests
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
ASSET = ROOT / "assets"
OUTPUT = ROOT / "vts_ocr_demo_antigravity.pptx"
TEMPLATE_URL = (
    "https://docs.google.com/presentation/d/"
    "1q_pOTwJijSkRhqOXbuSWdpNyFvmabVpO/export/pptx"
)
CACHE_TEMPLATE = Path("/tmp/vts_template_source.pptx")

RED = RGBColor(0xEE, 0x00, 0x33)
INK = RGBColor(0x44, 0x49, 0x4D)
GRAY = RGBColor(0x73, 0x7A, 0x80)
LIGHT = RGBColor(0xF4, 0xF4, 0xF4)
BORDER = RGBColor(0xE3, 0xE3, 0xE3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x20, 0x70, 0x55)
TITLE_FONT = "Sarabun"
BODY_FONT = "Arial"


def clone_shape(source, target, shape):
    """Copy a real VTS decorative shape, including image-fill relationships."""
    element = deepcopy(shape.element)
    for blip in element.xpath(".//a:blip"):
        old_id = blip.get(qn("r:embed"))
        if old_id:
            image_part = source.part.related_part(old_id)
            blip.set(qn("r:embed"), target.part.relate_to(image_part, RT.IMAGE))
    target.shapes._spTree.insert_element_before(element, "p:extLst")


def base_slide(prs, source, indices):
    slide = prs.slides.add_slide(source.slide_layout)
    for idx in indices:
        clone_shape(source, slide, source.shapes[idx])
    return slide


def textbox(slide, x, y, w, h, text, *, size=20, color=INK, bold=False,
            font=BODY_FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
            margin=0.02, spacing=1.12):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    for i, line in enumerate(text.split("\n")):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.text = line
        para.alignment = align
        para.space_after = Pt(size * 0.14)
        para.line_spacing = spacing
        for run in para.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return shape


def panel(slide, x, y, w, h, *, fill=LIGHT, edge=BORDER, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = edge
    shape.line.width = Pt(0.6)
    return shape


def card(slide, x, y, w, h, title, body, *, tint=LIGHT, emphasis=None,
         title_size=20, body_size=18):
    panel(slide, x, y, w, h, fill=tint)
    if emphasis:
        textbox(slide, x + 0.25, y + 0.2, w - 0.5, 0.62, emphasis,
                size=32, color=RED, bold=True, font=TITLE_FONT)
        start = y + 0.94
    else:
        start = y + 0.28
    textbox(slide, x + 0.25, start, w - 0.5, 0.72, title,
            size=title_size, bold=True, font=TITLE_FONT)
    textbox(slide, x + 0.25, start + 0.75, w - 0.5,
            max(0.3, h - (start - y) - 1.0), body, size=body_size)


def badge(slide, x, y, w, text, *, fill=LIGHT, color=INK):
    panel(slide, x, y, w, 0.38, fill=fill, edge=fill)
    textbox(slide, x + 0.07, y + 0.025, w - 0.14, 0.30, text,
            size=10, color=color, bold=True, align=PP_ALIGN.CENTER)


def photo(slide, filename, x, y, w, h, *, caption="", bg=WHITE):
    panel(slide, x, y, w, h, fill=bg, radius=False)
    path = ASSET / filename
    with Image.open(path) as im:
        iw, ih = im.size
        data = BytesIO()
        # PowerPoint and Google Slides do not support WebP images in PPTX containers.
        # Keep crops lossless PNG, encode full-page scans as high quality JPEG.
        is_page = filename.startswith(("demo-", "stamp-", "sign-", "missing-"))
        if is_page:
            im.convert("RGB").save(data, "JPEG", quality=92)
        else:
            im.save(data, "PNG")
        data.seek(0)
    image_h = h - (0.48 if caption else 0.12)
    scale = min((w - 0.15) / iw, (image_h - 0.13) / ih)
    pw, ph = iw * scale, ih * scale
    slide.shapes.add_picture(data, Inches(x + (w - pw) / 2),
                             Inches(y + (image_h - ph) / 2),
                             width=Inches(pw), height=Inches(ph))
    if caption:
        textbox(slide, x + 0.15, y + h - 0.42, w - 0.3, 0.34,
                caption, size=10, color=GRAY)


def quote(slide, x, y, w, label, value, *, good=False, h=0.95, size=17):
    fill = RGBColor(0xF2, 0xF7, 0xF4) if good else RGBColor(0xFA, 0xF0, 0xF0)
    panel(slide, x, y, w, h, fill=fill, edge=fill)
    textbox(slide, x + 0.17, y + 0.1, w - 0.34, 0.23, label.upper(),
            size=10, color=GREEN if good else RED, bold=True)
    textbox(slide, x + 0.17, y + 0.37, w - 0.34, h - 0.44, value,
            size=size, bold=True)


def banner(slide, y, text, *, size=21):
    panel(slide, 1.36, y, 17.4, 0.9,
          fill=RGBColor(0xFB, 0xF1, 0xF2), edge=RGBColor(0xF4, 0xD3, 0xD8))
    textbox(slide, 1.6, y + 0.15, 16.9, 0.6, text,
            size=size, color=INK, bold=True)


def table_slide(slide, x, y, w, h, headers, rows_data, col_widths=None):
    table_shape = slide.shapes.add_table(len(rows_data) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = table_shape.table
    if col_widths:
        for idx, cw in enumerate(col_widths):
            tbl.columns[idx].width = Inches(cw)

    for col_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x17, 0x22, 0x35)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.12)
        cell.margin_top = cell.margin_bottom = Inches(0.05)
        p = cell.text_frame.paragraphs[0]
        p.text = h_text
        p.alignment = PP_ALIGN.LEFT
        for r in p.runs:
            r.font.name = TITLE_FONT
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = WHITE

    for row_idx, row in enumerate(rows_data):
        is_required = 'required' in row[2].lower()
        is_out = 'out of scope' in row[2].lower()
        row_bg = RGBColor(0xFF, 0xFF, 0xFF) if row_idx % 2 == 0 else RGBColor(0xF6, 0xF8, 0xFA)
        if is_out:
            row_bg = RGBColor(0xFD, 0xFA, 0xFA)

        for col_idx, val in enumerate(row):
            cell = tbl.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = row_bg
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.alignment = PP_ALIGN.LEFT
            for r in p.runs:
                r.font.name = BODY_FONT
                r.font.size = Pt(11)
                r.font.color.rgb = GRAY if is_out else INK
                if col_idx == 0:
                    r.font.bold = True
                    if is_required:
                        r.font.color.rgb = RED
                elif col_idx == 1:
                    r.font.name = 'Consolas'
                    r.font.size = Pt(10.5)


def white_logo_png(template_bytes):
    from zipfile import ZipFile
    with ZipFile(BytesIO(template_bytes)) as archive:
        img = Image.open(BytesIO(archive.read("ppt/media/image3.png"))).convert("RGBA")
    pixels = img.load()
    for yy in range(img.height):
        for xx in range(img.width):
            r, g, b, a = pixels[xx, yy]
            if a:
                pixels[xx, yy] = (255, 255, 255, a)
    stream = BytesIO()
    img.save(stream, "PNG")
    stream.seek(0)
    return stream


def build_deck():
    if CACHE_TEMPLATE.exists():
        template_bytes = CACHE_TEMPLATE.read_bytes()
    elif Path("/tmp/opencode/vts_template_source.pptx").exists():
        template_bytes = Path("/tmp/opencode/vts_template_source.pptx").read_bytes()
    else:
        resp = requests.get(TEMPLATE_URL, timeout=90)
        resp.raise_for_status()
        template_bytes = resp.content

    if not template_bytes.startswith(b"PK"):
        raise ValueError("VTS template download is not a valid PPTX zip file")

    prs = Presentation(BytesIO(template_bytes))
    cover_source = prs.slides[0]
    content_source = prs.slides[3]
    closing_source = prs.slides[23]
    original_count = len(prs.slides)
    white_logo = white_logo_png(template_bytes)
    pages = []

    def content(title, source, section, *, number=None):
        print(f"Composing {len(pages) + 1:02d}: {title}", flush=True)
        slide = base_slide(prs, content_source, range(8))
        pages.append(slide)
        n = number or len(pages)
        textbox(slide, 1.33, 0.59, 15.5, 0.68, title,
                size=28, color=INK, bold=True, font=TITLE_FONT)
        textbox(slide, 1.35, 10.27, 14.3, 0.35, source,
                size=9, color=GRAY)
        textbox(slide, 18.78, 10.27, 0.65, 0.35, str(n).zfill(2),
                size=10, color=RED, align=PP_ALIGN.RIGHT)
        badge(slide, 1.36, 1.43, min(5.5, 1.2 + len(section) * 0.11), section.upper(),
              fill=RGBColor(0xF8, 0xE7, 0xEA), color=RED)
        return slide

    # 01 — Cover slide
    print("Composing 01: Cover", flush=True)
    s = base_slide(prs, cover_source, range(6))
    pages.append(s)
    s.shapes.add_picture(white_logo, Inches(16.25), Inches(0.75), width=Inches(3.0))
    textbox(s, 8.05, 3.9, 10.8, 1.7, "AI OCR cho văn bản\nhành chính",
            size=44, color=WHITE, bold=True, font=TITLE_FONT)
    textbox(s, 8.08, 6.15, 10.6, 1.1,
            "Demo pipeline 01 · Văn phòng số / VPS", size=23, color=WHITE)
    textbox(s, 8.08, 7.43, 10.6, 0.65,
            "Báo cáo 29/09/2026  |  Full inference: 35/100 tài liệu đạt",
            size=16, color=WHITE)

    # 02 — Bối cảnh
    s = content("Từ nhập liệu thủ công đến dữ liệu có cấu trúc",
                "Nguồn: outline báo cáo · mục 1", "01 · Bối cảnh")
    card(s, 1.35, 2.06, 5.5, 5.7, "Hiện trạng",
         "Nhân viên đọc PDF rồi nhập lại số ký hiệu, ngày ban hành, cơ quan, trích yếu… vào hệ thống.", emphasis="01")
    card(s, 7.07, 2.06, 5.5, 5.7, "Pain point",
         "Khối lượng xử lý lớn làm tăng thời gian, nhân lực kiểm tra và nguy cơ sai sót; khó mở rộng theo nhu cầu.", emphasis="02")
    card(s, 12.78, 2.06, 5.9, 5.7, "Mục tiêu PoC",
         "Kiểm chứng giải pháp OCR/KIE in-house có thể tùy chỉnh theo nghiệp vụ và từng bước thử nghiệm tại UAT VPS.", emphasis="03")
    banner(s, 8.15, "Đây là bước đánh giá khả thi, chưa phải cam kết thay thế dịch vụ hiện tại.")

    # 03 — Agile
    s = content("Viettel Agile: học từ dữ liệu thật",
                "Nguồn: outline báo cáo · mục 1; chuẩn benchmark §6, §9", "01 · Bối cảnh")
    steps = [("Aware", "Nhận diện pain point"), ("Go Through", "Khảo sát nghiệp vụ"),
             ("Implement", "Xây baseline"), ("Learn", "Benchmark & đo lỗi"),
             ("Enable", "Hướng tới áp dụng")]
    for i, (step, detail) in enumerate(steps):
        x = 1.36 + i * 3.53
        card(s, x, 2.32, 3.3, 2.95, step, detail,
             tint=RGBColor(0xFA, 0xF1, 0xF3) if i == 3 else LIGHT,
             title_size=20, body_size=16)
    card(s, 1.36, 5.72, 8.45, 3.38, "Nhu cầu khách hàng là điểm xuất phát",
         "Đối chiếu biểu mẫu với BA để phân biệt trường nào OCR đọc từ PDF, trường nào người dùng nhập, trường nào hệ thống tự điền.",
         body_size=19)
    card(s, 10.1, 5.72, 8.67, 3.38, "Tối ưu có căn cứ",
         "Mọi thay đổi cần GT, benchmark, trace và regression test; không sửa luật chỉ vì một tên file hoặc một đáp án GT cụ thể.",
         body_size=19)

    # 04 — Phạm vi dataset
    s = content("Benchmark 100 PDF — đủ đa dạng để lộ điểm yếu",
                "Nguồn: full inference manifest.json + predictions.jsonl · 100 PDF", "02 · Phạm vi")
    for i, (value, title, detail) in enumerate([
        ("100", "văn bản", "Có GT gắn theo tài liệu"),
        ("2.249", "trang", "1–353 trang / tài liệu"),
        ("69", "scan", "Chất lượng OCR biến thiên"),
        ("31", "điện tử", "Print / digital PDF"),
    ]):
        card(s, 1.36 + 4.42 * i, 2.02, 4.13, 3.35, title, detail,
             emphasis=value, title_size=20, body_size=16)
    card(s, 1.36, 5.85, 8.45, 3.38, "Loại văn bản",
         "Quyết định (30), Công văn (17), Nghị quyết, Thông báo, Thông tư, Kế hoạch…\nĐộ dài và cấu trúc bố cục đa dạng.", body_size=19)
    card(s, 10.1, 5.85, 8.67, 3.38, "Nguồn và chủ thể",
         "Cơ quan Chính phủ, bộ/ngành, địa phương, sở giao dịch chứng khoán, doanh nghiệp… Có scan, dấu mộc, chữ viết tay; chưa đủ phủ lớp hiếm.", body_size=19)

    # 05 — Hợp đồng API & cách chấm
    s = content("Định nghĩa đầu ra trước khi đo độ chính xác",
                "Nguồn: docs/ocr_pipeline_benchmark_metrics.md · §1, §5", "02 · Phạm vi")
    card(s, 1.36, 2.03, 8.45, 5.66, "7 trường bắt buộc & 5 trường đánh giá thêm",
         "7 trường required:\ntype · title · code · documentDate\nofficeSender · priority_level · security_level\n\n5 trường optional được chấm:\nrecipients · signer · first_recipients\nsigner_title · province",
         body_size=19)
    card(s, 10.1, 2.03, 8.67, 5.66, "13 key trong API, chỉ chấm 12 field",
         "Theo chuẩn hiện hành, receiverDate giữ trong schema nhưng phải trả chuỗi rỗng và không chấm. Bundle full inference lịch sử còn lưu giá trị legacy; bản báo cáo loại field này.\n\nDocument Success@5%: mọi required field có GT trong cùng một tài liệu đều phải đạt Exact hoặc Acceptable theo metric tương ứng.",
         tint=RGBColor(0xFA, 0xF1, 0xF3), body_size=18)
    banner(s, 8.02, "Phân bucket: Exact / Acceptable / Wrong / Missing / Hallucination / Correct Absent.", size=20)

    # 06 — Bảng danh mục 13 trường API & tiêu chí đánh giá
    s = content("Danh mục 13 trường API & tiêu chí đánh giá",
                "Nguồn: docs/ocr_pipeline_benchmark_metrics.md · §1, §5", "02 · Phạm vi")
    headers = ["Mô tả ngắn tiếng Việt", "Key API", "Nhóm field", "Metric đánh giá", "Target F1 / NEM"]
    field_rows = [
        ["Loại văn bản", "type", "Critical + required", "F1 Exact (khớp tuyệt đối)", "F1 ≥ 0.85 (NEM ≥ 0.75)"],
        ["Trích yếu / Tiêu đề", "title", "Critical + required", "F1@5% (CER ≤ 5%)", "F1 ≥ 0.85 (NEM ≥ 0.75)"],
        ["Số, ký hiệu văn bản", "code", "Critical + required", "F1 Exact (chuẩn hóa gạch nối)", "F1 ≥ 0.85 (NEM ≥ 0.75)"],
        ["Ngày ban hành", "documentDate", "Critical + required", "F1 Exact (chuẩn dd/mm/yyyy)", "F1 ≥ 0.85 (NEM ≥ 0.75)"],
        ["Cơ quan ban hành", "officeSender", "Critical + required", "F1@5% (CER ≤ 5%)", "F1 ≥ 0.85 (NEM ≥ 0.75)"],
        ["Độ khẩn", "priority_level", "Critical + required", "F1 Exact (Enum 4 mức)", "F1 ≥ 0.85 (NEM ≥ 0.75)"],
        ["Độ mật", "security_level", "Critical + required", "F1 Exact (Enum 4 mức)", "F1 ≥ 0.85 (NEM ≥ 0.75)"],
        ["Nơi nhận cuối VB", "recipients", "Critical + optional", "F1@5% (CER ≤ 5%)", "F1 ≥ 0.85 (NEM ≥ 0.75)"],
        ["Họ tên người ký", "signer", "Critical + optional", "F1 Exact (khớp tên)", "F1 ≥ 0.85 (NEM ≥ 0.75)"],
        ["Kính gửi đầu VB", "first_recipients", "Critical + optional", "F1@5% (CER ≤ 5%)", "F1 ≥ 0.85 (NEM ≥ 0.75)"],
        ["Chức danh người ký", "signer_title", "Critical + optional", "F1@5% (CER ≤ 5%)", "F1 ≥ 0.85 (NEM ≥ 0.75)"],
        ["Địa danh ban hành", "province", "Non-critical + optional", "F1 Exact (khớp địa danh)", "F1 ≥ 0.85 (NEM ≥ 0.75)"],
        ["Ngày tiếp nhận", "receiverDate", "Out of scope", "Không chấm (luôn trả \"\")", "— (Ngoài phạm vi)"],
    ]
    table_slide(s, 1.36, 2.05, 17.4, 7.8, headers, field_rows, [3.8, 2.3, 3.3, 4.1, 3.9])

    # 07 — Pipeline luồng xử lý
    s = content("Team đã xây gì, và pipeline chạy ra sao?",
                "Nguồn: outline báo cáo · mục 2; chuẩn benchmark §6–§8", "02 · Phạm vi")
    for i, text in enumerate(["PDF", "Layout + OCR", "Extraction", "Chuẩn hóa API", "Benchmark + Evidence"]):
        panel(s, 1.36 + i * 3.52, 2.26, 3.27, 1.45, fill=LIGHT)
        textbox(s, 1.54 + i * 3.52, 2.65, 2.93, 0.65, text,
                size=20, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
    for i, (title, body) in enumerate([
        ("Chuẩn bị", "Thu thập PDF; xây GT; chốt quy tắc nghiệp vụ và hợp đồng 13-field API với BA."),
        ("Thử nghiệm", "Các thành viên khảo sát pipeline riêng; chọn pipeline 01 làm baseline để demo."),
        ("Đo & truy vết", "Chấm chất lượng, thời gian, RAM/GPU; giữ ảnh trang, bbox và OCR nguồn để soi lỗi."),
    ]):
        card(s, 1.36 + i * 5.83, 4.38, 5.55, 3.5, title, body, body_size=18)
    banner(s, 8.3, "Số liệu trình bày lấy từ full inference (PDF → model → extraction), không trộn tài nguyên của replay.")

    # 07 — Demo overview
    s = content("Demo theo 3 mức độ: dễ → phổ biến → khó",
                "Nguồn: full inference predictions.jsonl, document_success.csv, field_errors.csv", "03 · Demo")
    card(s, 1.36, 2.08, 5.5, 5.58, "Case dễ",
         "3116+QĐ-BKHCN và 26/2026/NQ-HĐND:\nĐầu ra sạch, không có field Wrong/Missing/Hallucination.",
         emphasis="01", body_size=19)
    card(s, 7.08, 2.08, 5.5, 5.58, "Case phổ biến",
         "Quyết định 1666 sạch 12 field;\nCông văn 8906 đạt đủ 7 required nhưng miss 2 field optional.",
         emphasis="02", body_size=19)
    card(s, 12.8, 2.08, 5.98, 5.58, "Case khó",
         "8931, 8948, 29+CT-TTg: scan;\nlỗi tập trung ở ngày ban hành, dấu khẩn, chữ ký và chức danh.",
         emphasis="03", body_size=19)
    banner(s, 8.12, "Mỗi case: PDF đầu vào → 13 key đầu ra → GT/prediction → ảnh + trace → thời gian xử lý.")

    # 08 — Easy case
    s = content("Văn bản rõ: trích xuất đúng và truy vết được",
                "Nguồn: predictions.jsonl + field_errors.csv · full inference", "03 · Demo / Case dễ")
    photo(s, "demo-3116.webp", 1.37, 2.02, 7.15, 7.7, caption="3116+QĐ-BKHCN · trace trang 1 full inference")
    card(s, 8.86, 2.02, 9.9, 3.54, "3116+QĐ-BKHCN · 2 trang",
         "4,37 s/file · Document Success PASS. Trong 12 field được chấm theo policy hiện hành: 10 Exact, title Acceptable, first_recipients Correct Absent.",
         body_size=19)
    card(s, 8.86, 5.82, 9.9, 3.42, "26+2026+NQ-HĐND · 5 trang",
         "7,93 s/file · Document Success PASS; không có cả bucket Acceptable, Wrong, Missing hoặc Hallucination.", body_size=19)

    # 09 — Common case
    s = content("2 dạng chiếm 47% tập mẫu: Quyết định và Công văn",
                "Nguồn: predictions.jsonl + document_success.csv + field_errors.csv", "03 · Demo / Phổ biến")
    photo(s, "demo-1666.webp", 1.36, 2.07, 5.6, 6.75, caption="1666+QĐ-TTg · scan 5 trang")
    photo(s, "demo-8906.webp", 7.16, 2.07, 5.6, 6.75, caption="8906+VPCP-CN · print 2 trang")
    card(s, 13.02, 2.07, 5.75, 3.16, "Quyết định 1666",
         "12,52 s/file · PASS;\nKhông field nào Wrong/Missing/Hallucination.", body_size=17)
    card(s, 13.02, 5.49, 5.75, 3.33, "Công văn 8906",
         "4,89 s/file · 7/7 required đạt.\nrecipients và first_recipients Missing.", body_size=17)

    # 10 — Difficult case overview
    s = content("Header đúng chưa đồng nghĩa đúng cả văn bản",
                "Nguồn: full inference document_success.csv + field_errors.csv + trace v2", "03 · Demo / Khó")
    card(s, 1.36, 2.16, 5.52, 5.72, "8931+VPCP-PL",
         "4,42 s · 1 trang scan.\nNgày 29→19; cơ quan dính chữ; Hỏa tốc thành Bình thường.\nFAIL 3 required.",
         body_size=19)
    card(s, 7.1, 2.16, 5.52, 5.72, "8948+VPCP-CN",
         "6,50 s · 2 trang scan.\nNgày Missing, Kính gửi Missing; vùng chức danh lẫn OCR rác.\nFAIL vì ngày.",
         body_size=19)
    card(s, 12.84, 2.16, 5.93, 5.72, "29+CT-TTg",
         "7,76 s · 4 trang scan.\nHeader chính Exact; rớt Hỏa tốc, tên người ký sai, chức danh Missing.\nFAIL vì độ khẩn.",
         body_size=19)
    banner(s, 8.32, "Các slide lỗi sau đây gắn GT/prediction với đúng ảnh và OCR trong trace; không suy đoán vị trí GT.")

    # 11 — Results & Provenance
    s = content("35/100 tài liệu đạt; PoC còn cách 45 tài liệu",
                "Nguồn: summary.json · full inference v3 · metric v2.1 · Quadro RTX 5000", "04 · Kết quả")
    card(s, 1.36, 2.03, 5.55, 3.22, "Document Success@5%",
         "35 PASS / 100 văn bản. PoC ≥80%; stretch ≥90%.", emphasis="35%", body_size=18)
    card(s, 7.13, 2.03, 5.55, 3.22, "Thời gian trung bình",
         "44,1 s / văn bản (wall-clock, 100 PDF; số trang biến thiên lớn).", emphasis="44,1 s", body_size=18)
    card(s, 12.9, 2.03, 5.88, 3.22, "Giây / trang*",
         "1,14–5,85 s; *TB của 100 tỷ số (giây/file ÷ trang/file).", emphasis="2,48 s", body_size=18)
    card(s, 1.36, 5.62, 8.45, 3.53, "RAM tiến trình",
         "Trung bình 2.745 MB · đỉnh 6.012 MB.\nBenchmark chạy trên GPU Quadro RTX 5000.", body_size=19)
    card(s, 10.1, 5.62, 8.68, 3.53, "VRAM: phân biệt phạm vi đo",
         "Device VRAM toàn GPU: 3.077 / 6.793 MB (TB/đỉnh).\nProcess VRAM: chưa đo được; không gán 0 hay thay bằng allocator.", body_size=18)

    # 12 — Field results table (compact)
    s = content("Điểm nghẽn: chức danh ký, trích yếu, ngày",
                "Nguồn: full inference metrics.csv · receiverDate ngoài phạm vi chấm hiện hành", "04 · Kết quả theo field")
    card(s, 1.36, 2.04, 8.46, 6.92, "Tín hiệu tích cực",
         "province: F1 = 0,984 — field duy nhất đạt đủ target đánh giá.\n\ntype, security_level: F1 = 1,00 nhưng chưa đánh giá đầy đủ lớp hiếm / mẫu GT trống.\n\nKhông có Hallucination trong các field có GT trống ở bundle này; không suy rộng sang lớp chưa phủ.",
         body_size=20)
    card(s, 10.1, 2.04, 8.68, 6.92, "Field cần ưu tiên (Bảng F1)",
         "signer_title     F1 0,380  ·  51 Wrong + 12 Missing\n\ntitle            F1 0,670  ·  31 Wrong + 3 Missing\n\nofficeSender     F1 0,735  ·  24 Wrong + 4 Missing\n\ndocumentDate     F1 0,785  ·  10 Wrong + 19 Missing\n\n(Bảng chỉ số field gọn, không vẽ biểu đồ lỗi)",
         tint=RGBColor(0xFA, 0xF1, 0xF3), body_size=18)

    # 13 — Evidence OCR
    s = content("Một ký tự OCR có thể làm hỏng trường bắt buộc",
                "Nguồn: field_errors.csv + trace v2 (8931+VPCP-PL) trang 1 · full inference",
                "04 · Evidence / OCR")
    photo(s, "ocr-date-8931.webp", 1.36, 2.14, 7.65, 2.2,
          caption="Ngày ban hành · bbox [847, 274, 1337, 308]")
    photo(s, "ocr-sender-8931.webp", 1.36, 4.7, 7.65, 2.2,
          caption="Cơ quan ban hành · bbox [236, 154, 663, 195]")
    quote(s, 9.35, 2.14, 9.43, "GT · documentDate", "29/08/2026", good=True)
    quote(s, 9.35, 3.31, 9.43, "Pred / OCR", "19/08/2026 · ngày 19 tháng 8 năm 2026")
    quote(s, 9.35, 4.67, 9.43, "GT · officeSender", "VĂN PHÒNG CHÍNH PHỦ", good=True)
    quote(s, 9.35, 5.84, 9.43, "Pred / OCR", "VĂN PHÒNG CHÍNH PHỦNH PHỦ")
    textbox(s, 9.37, 7.13, 9.32, 2.2,
            "Nhận định: Extraction bám theo OCR nguồn. Metadata có chữ viết tay nhưng chưa đủ bằng chứng để quy lỗi cụ thể cho chữ viết tay ở hai vùng này.",
            size=18)

    # 14 — Evidence Stamp / Sign
    s = content("Header đúng, nhưng mức khẩn và cụm ký vẫn rớt",
                "Nguồn: field_errors.csv + trace v2 (29+CT-TTg) trang 1 & 4 · full inference",
                "04 · Evidence / Dấu & ký")
    photo(s, "stamp-29.webp", 1.36, 2.07, 4.98, 7.52,
          caption="Trang 1 · kiểm tra vùng dấu Hỏa tốc (pred không có evidence)")
    photo(s, "sign-29.webp", 6.54, 2.07, 4.98, 7.52,
          caption="Trang 4 · vùng ký, OCR đọc 'Huốc Dũng'")
    quote(s, 11.79, 2.07, 6.98, "GT / Pred · priority_level",
          "1_HỎA TỐC → 0_BÌNH THƯỜNG", h=1.35, size=16)
    quote(s, 11.79, 3.62, 6.98, "GT / Pred · signer",
          "Hồ Quốc Dũng → Huốc Dũng", h=1.35, size=16)
    quote(s, 11.79, 5.17, 6.98, "GT / Pred · signer_title",
          "KT. THỦ TƯỚNG PHÓ THỦ TƯỚNG → “ ”", h=1.38, size=14)
    textbox(s, 11.85, 6.88, 6.9, 2.4,
            "Nhận định: Vùng dấu/chữ ký cần review ảnh màu; priority_level không có evidence. Chưa khẳng định dấu gây ra toàn bộ lỗi vùng ký nếu không khoanh trực tiếp được vùng bị đè.", size=17)

    # 15 — Evidence Layout / Crop
    s = content("Vùng text lẫn nội dung: trích yếu và cơ quan bị kéo sai",
                "Nguồn: field_errors.csv + trace v2 (12+NQ-PVIRe & 240+TB-PVIRe) · full inference",
                "04 · Evidence / Layout")
    photo(s, "layout-title-12.webp", 1.36, 2.13, 8.42, 4.05,
          caption="12+NQ · khung đỏ là bbox prediction evidence text-after-document-type")
    photo(s, "layout-sender-240.webp", 10.1, 2.13, 8.68, 4.05,
          caption="240+TB · khung đỏ là bbox prediction evidence upper-left-organization ('Số:')")
    card(s, 1.36, 6.45, 8.42, 2.73, "12+NQ-PVIRe · title Wrong",
         "GT: “Về việc chốt danh sách cổ đông…”\nPred: thêm “SỞ GIAO DỊCH CHỨNG KHOÁN HÀ NỘI … HỘI ĐỒNG QUẢN TRỊ”.",
         body_size=16)
    card(s, 10.1, 6.45, 8.68, 2.73, "240+TB-PVIRe · officeSender Wrong",
         "GT: “TỔNG CÔNG TY CỔ PHẦN TÁI BẢO HIỂM PVI”\nPred: “Số:” — bắt nhầm anchor làm cơ quan ban hành.", body_size=16)

    # 16 — Evidence Extraction / Heuristics
    s = content("OCR có text, hậu xử lý vẫn giữ nhầm một “stub”",
                "Nguồn: field_errors.csv + trace v2 (2125+TB-ĐLTKV) · rule kinh-gui-anchor",
                "04 · Evidence / Extraction")
    photo(s, "extract-2125.webp", 1.36, 2.14, 17.4, 2.83,
          caption="Vùng Kính gửi · bbox [473, 616, 1450, 767]")
    quote(s, 1.36, 5.25, 17.4, "GT · first_recipients",
          "Ủy ban Chứng khoán Nhà nước; Sở Giao dịch Chứng khoán Hà Nội; Trung tâm Lưu ký Chứng khoán Việt Nam.",
          good=True, h=1.1, size=17)
    quote(s, 1.36, 6.55, 17.4, "Prediction",
          "Uỷ ban Chứng khoán Nhà nước; số 1.772; Sở Giao dịch Chứng khoán Hà Nội; Trung tâm Lưu ký Chứng khoán Việt Nam.",
          h=1.15, size=17)
    textbox(s, 1.45, 8.12, 17.05, 1.06,
            "Nhận định: Nguồn OCR chứa 'số 1.772'. Lỗi đáng thử ở bước lọc stub khi tách danh sách; cần regression test để tránh xóa số thật trong tên người nhận.",
            size=18)

    # 17 — Evidence Missing
    s = content("Không thấy evidence ≠ GT ở vị trí mình đoán",
                "Nguồn: field_errors.csv + trace v2 (8948+VPCP-CN) · không suy diễn vị trí GT",
                "04 · Evidence / Missing")
    photo(s, "missing-8948.webp", 1.36, 2.06, 7.12, 7.53,
          caption="Trang 1 do người review chọn để kiểm tra; không phải bbox GT")
    quote(s, 8.78, 2.06, 10, "GT / Pred · documentDate",
          "01/09/2026 → “ ” · Missing · không có evidence", h=1.1, size=18)
    quote(s, 8.78, 3.42, 10, "GT / Pred · first_recipients",
          "Bộ trưởng các Bộ…; Chủ tịch UBND… → “ ” · Missing · không có evidence", h=1.35, size=16)
    card(s, 8.78, 5.06, 10, 4.02, "Nhận định có giới hạn",
         "Cần soi layout/OCR trang đầu và ranh giới văn bản trước khi quyết định fallback. Pipeline không bịa giá trị ngày khi không đọc được.\n\nBundle không có Hallucination ở field có GT trống; đây là tín hiệu tốt, nhưng không suy rộng sang lớp hiếm.",
         body_size=18)

    # 18 — Evidence GT Review
    s = content("Không phải mọi bucket Wrong đều là OCR đọc sai",
                "Nguồn: field_errors.csv + trace v2 · 03+BC.HĐQT & 27+2026+VBHN",
                "04 · Evidence / GT")
    photo(s, "gt-03.webp", 1.36, 2.1, 8.43, 3.35,
          caption="OCR nguồn: “Số 03/BC HĐQT” · trang 1")
    photo(s, "gt-27.webp", 10.1, 2.1, 8.68, 3.35,
          caption="OCR nguồn: “Số 27/2026/VBHN-NĐ-BTC” · trang 1")
    card(s, 1.36, 5.7, 8.43, 3.58, "03+BC.HĐQT",
         "GT: 03/BC.HĐQT · Pred: 03/BC-HĐQT.\nQuy tắc API biến dấu cách còn lại thành '-'; cần xác nhận cách gán GT đối với dấu chấm '.'",
         body_size=18)
    card(s, 10.1, 5.7, 8.68, 3.58, "27+2026+VBHN",
         "GT: …-ND-BTC · Pred: …-NĐ-BTC.\nẢnh/trace ủng hộ chữ 'Đ'; cần reviewer xác nhận, không tự ý sửa GT.",
         body_size=18)

    # 19 — Phân bố nguyên nhân lỗi (Error distribution)
    s = content("229 ô field lỗi — phân loại sơ bộ 4 nhóm",
                "Nguồn: full inference v3 field_errors.csv + trace + error_cause_full_inference_v3.csv",
                "04 · Phân bố lỗi")
    for i, (pct, cnt, grp_title, detail, tint) in enumerate([
        ("47,6%", "109 ô", "OCR / chính tả",
         "Nhận dạng sai ký tự, rác scan, dính chữ, đọc nhầm số/chữ (VD: ngày 29→19, Văn phòng Chính phủnh phủ, Huốc Dũng).",
         RGBColor(0xFA, 0xF0, 0xF0)),
        ("17,9%", "41 ô", "Trích xuất & Layout",
         "Bắt nhầm phân vùng, kéo dài trích yếu lẫn tên cơ quan, giữ nhầm stub (2125), bắt nhầm anchor 'Số:' (240).",
         LIGHT),
        ("2,2%", "5 ô", "Dấu mộc & Vùng ký",
         "Dấu đỏ Hỏa tốc bị miss (29+CT), con dấu đỏ đè làm sai lệch chữ ký hoặc mất chức danh người ký.",
         LIGHT),
        ("32,3%", "74 ô", "Khác & Cần review",
         "Missing không có visual evidence (8948), nghi vấn quy tắc/chính sách gán nhãn GT (03/BC, 27/VBHN), cần rà soát thêm.",
         LIGHT),
    ]):
        x = 1.36 + i * 4.42
        card(s, x, 2.05, 4.13, 5.85, f"{grp_title} ({cnt})", detail,
             tint=tint, emphasis=pct, title_size=18, body_size=15)
    banner(s, 8.22, "Mẫu số: 229 ô lỗi (156 Wrong + 73 Missing trên 12 field). OCR nguồn là điểm nghẽn lớn nhất (47,6%).", size=19)

    # 20 — Feasibility & Gate
    s = content("Khả thi để tiếp tục R&D; chưa đủ điều kiện thay thế",
                "Nguồn: outline báo cáo · mục 5; kết quả full inference v3", "05 · Đánh giá")
    card(s, 1.36, 2.12, 8.44, 5.75, "Đã chứng minh",
         "• Pipeline chạy end-to-end trên 100 PDF, có API dữ liệu cấu trúc.\n\n• Có GT, benchmark version hóa, ảnh trang, OCR và trace để lặp lại phân tích.\n\n• Có case print/scan hoạt động tốt; có thể ưu tiên cải thiện theo từng trường.",
         body_size=19)
    card(s, 10.1, 2.12, 8.68, 5.75, "Chưa chứng minh",
         "• 35% Document Success còn cách target PoC 80%.\n\n• Mẫu lớp mật/khẩn hiếm và tài liệu chụp nghiêng/ngang chưa đủ coverage.\n\n• Chưa có đo chi phí và benchmark đối chứng đủ điều kiện để kết luận thay thế API hiện tại.",
         tint=RGBColor(0xFA, 0xF1, 0xF3), body_size=19)
    banner(s, 8.17, "Cân bằng tổng quát hóa vs. overfit; chất lượng vs. GPU; hiểu ngữ cảnh vs. thời gian.", size=20)

    # 21 — Next steps & Roadmap
    s = content("Ưu tiên accuracy, sau đó mới tối ưu vận hành",
                "Nguồn: outline báo cáo · mục 6; chuẩn benchmark §6, §9", "06 · Tiếp theo")
    card(s, 1.36, 2.05, 5.5, 5.82, "01 · Dữ liệu",
         "Bổ sung hard cases: dấu đỏ, chữ viết tay, trang nghiêng/ngang, lớp mật/khẩn và văn bản dài. Rà soát GT theo policy đã chốt; tách train/dev/test nếu fine-tune.",
         body_size=18)
    card(s, 7.08, 2.05, 5.5, 5.82, "02 · Kỹ thuật",
         "Benchmark riêng OCR; thử model chữ viết tay và xử lý dấu/watermark; cải thiện ranh giới văn bản/trang cuối; chuẩn hóa tập giá trị hữu hạn khi có evidence.",
         body_size=18)
    card(s, 12.8, 2.05, 5.98, 5.82, "03 · Kiểm chứng",
         "Regression test sau mỗi thay đổi.\nĐạt accuracy target rồi mới tối ưu tốc độ → tài nguyên → chi phí → pilot/UAT có kiểm soát.",
         body_size=18)
    banner(s, 8.2, "Cần thống nhất mốc review kế tiếp, bộ hard cases ưu tiên và tài nguyên GPU/chuyên môn.", size=20)

    # 22 — Closing
    s = base_slide(prs, closing_source, range(6))
    pages.append(s)
    white_logo.seek(0)
    s.shapes.add_picture(white_logo, Inches(16.2), Inches(0.7), width=Inches(3.0))
    textbox(s, 7.77, 4.23, 11.1, 1.25, "Xin cảm ơn!",
            size=50, color=WHITE, bold=True, font=TITLE_FONT)
    textbox(s, 7.8, 5.68, 11.05, 1.5,
            "35/100 hôm nay · 80/100 là mốc PoC\nXin feedback: thứ tự field ưu tiên, mức sai sót chấp nhận được và điều kiện pilot tại UAT.",
            size=23, color=WHITE)

    # Remove template source slides, preserving master styles & relationships
    for _ in range(original_count):
        slide_id = prs.slides._sldIdLst[0]
        prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)

    assert len(prs.slides) == len(pages) == 23
    prs.core_properties.title = "Demo AI OCR văn bản hành chính · Viettel Solutions"
    prs.core_properties.subject = "Pipeline 01 · full inference 35/100 · 29/09/2026"
    prs.save(OUTPUT)
    print(f"Successfully generated {OUTPUT} ({OUTPUT.stat().st_size:,} bytes; {len(prs.slides)} slides)")


if __name__ == "__main__":
    build_deck()
